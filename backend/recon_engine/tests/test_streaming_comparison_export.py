"""Regression coverage for the Auto/streaming-mode results pipeline.

Reproduces exactly what a completed Auto-mode run (created via the chatbot or
the auto wizard) looks like on disk — a `ReconciliationRun` with the
`streaming:<graph_run_id>` sentinel source/target snapshot ids and no
`shadow_id` at all (see `auto_pipeline.nodes._do_finalize`), plus per-batch
resolved value mappings that never get written back to `contract_store` (see
`auto_pipeline.nodes._do_run_batches`) — without going through the full
LangGraph + LLM/connector machinery. Before the fix, `build_comparison_workbook`
tried to re-join this run's field values from those sentinel/missing ids and
came back with every business-key/compare-field column blank, and an empty
Mapping Details sheet, exactly as reported against a real autorun export.
"""

from __future__ import annotations

from io import BytesIO

import openpyxl
import pandas as pd

from backend.recon_engine import service
from backend.recon_engine.compiler import StubContractCompiler
from backend.recon_engine.engine.executor import build_shadow_source
from backend.recon_engine.engine.reconciler import reconcile
from backend.recon_engine.models.run import ReconciliationRun, RunStatus
from backend.recon_engine.models.snapshot import RawLayer
from backend.recon_engine.models.value_mapping import ValueMapping
from backend.recon_engine.insights import builder as insights_builder
from backend.recon_engine.storage import result_store, run_store, run_value_mapping_store


def _streaming_run() -> str:
    graph_run_id = "autorun_test_streaming_fix"

    source_df = pd.DataFrame({
        "Material": ["A", "B"],
        "Plant": ["P1", "P1"],
        "Date": ["2024-01-01", "2024-01-02"],
        "ReqQty": [10, 20],
    })
    target_df = pd.DataFrame({
        "PRDID": ["PA", "B"],
        "LOCID": ["LOC1", "LOC1"],
        "Date": ["2024-01-01", "2024-01-02"],
        "SalesOrderRequest": [10, 25],
    })

    src = service.ingest_snapshot(source_df, layer=RawLayer.SOURCE, source_type="excel")
    tgt = service.ingest_snapshot(target_df, layer=RawLayer.TARGET, source_type="excel")

    # The approved contract carries NO value_mappings — an Auto-mode run
    # resolves those independently per date-batch (see
    # auto_pipeline.nodes._do_run_batches) rather than up front, same as in
    # production.
    draft, _ = service.compile_draft(
        mapping_sheet=[
            {"source_col": "Material", "target_col": "PRDID", "role": "key"},
            {"source_col": "Plant", "target_col": "LOCID", "role": "key"},
            {"source_col": "Date", "target_col": "Date", "role": "key"},
            {"source_col": "ReqQty", "target_col": "SalesOrderRequest", "role": "compare"},
        ],
        rules="",
        business_key=[
            {"source_field": "Material", "target_field": "PRDID"},
            {"source_field": "Plant", "target_field": "LOCID"},
            {"source_field": "Date", "target_field": "Date"},
        ],
        compare_fields=[{"source_field": "ReqQty", "target_field": "SalesOrderRequest"}],
        value_mappings=[],
        source_schema=["Material", "Plant", "Date", "ReqQty"],
        target_schema=["PRDID", "LOCID", "Date", "SalesOrderRequest"],
        comparison_type="c", source_type="excel", target_type="excel",
        compiler=StubContractCompiler(),
    )
    contract = service.approve_contract(draft, approved_by="alice")

    # One date-batch's resolved product/location pairing, exactly the shape
    # `_pair_field_for_batch` returns — folded into a per-batch contract copy
    # the way `_do_run_batches` does, never persisted to `contract_store`.
    product_mapping = ValueMapping.model_validate({
        "source_field": "Material", "target_field": "PRDID",
        "matches": [
            {"source_value": "A", "target_value": "PA", "confidence": "high", "rule": "t", "evidence": "e",
             "row_count": 1},
            {"source_value": "B", "target_value": "B", "confidence": "very_high", "rule": "t", "evidence": "e",
             "row_count": 1},
        ],
    })
    location_mapping = ValueMapping.model_validate({
        "source_field": "Plant", "target_field": "LOCID",
        "matches": [
            {"source_value": "P1", "target_value": "LOC1", "confidence": "high", "rule": "t", "evidence": "e",
             "row_count": 2},
        ],
    })
    run_value_mapping_store.record_batch_mapping(graph_run_id, product_mapping)
    run_value_mapping_store.record_batch_mapping(graph_run_id, location_mapping)

    batch_contract = contract.model_copy(update={"value_mappings": [product_mapping, location_mapping]})
    built = build_shadow_source(batch_contract, source_df)
    recon = reconcile(batch_contract, built.shadow_df, target_df)
    recon.detail_df = service.attach_field_values(
        recon.detail_df, batch_contract, shadow_df=built.shadow_df, target_df=target_df, raw_source_df=source_df,
    )

    detail_df = recon.detail_df.copy()
    detail_df["run_id"] = graph_run_id
    detail_df["batch_id"] = "batch_1"
    detail_df["record_id"] = [f"rec_{i}" for i in range(len(detail_df))]

    result = result_store.start_streaming_result(
        run_id=graph_run_id, contract_id=contract.contract_id, contract_version=contract.contract_version,
    )
    result_store.append_batch_result(result.result_id, detail_df=detail_df, batch_summary=recon.summary)

    # Mirrors auto_pipeline.nodes._do_finalize's ReconciliationRun exactly:
    # sentinel source/target snapshot ids, no shadow_id.
    run = ReconciliationRun(
        run_id=graph_run_id,
        contract_id=contract.contract_id,
        contract_version=contract.contract_version,
        source_snapshot_id=f"streaming:{graph_run_id}",
        target_snapshot_id=f"streaming:{graph_run_id}",
        status=RunStatus.COMPLETED,
        created_by="chat",
    )
    run_store.save_run(run)
    return graph_run_id


def test_streaming_run_all_records_has_real_field_values_not_blank():
    run_id = _streaming_run()
    content = service.build_comparison_workbook(run_id)
    wb = openpyxl.load_workbook(BytesIO(content))

    ws = wb["All Records"]
    header = next(ws.iter_rows(values_only=True))
    # The run-scoped accumulated value mappings make this an Original/Paired
    # split, exactly like a Manual-mode run whose contract carries the
    # mapping directly — not the flat "Material"/"Plant" single columns a
    # missing/empty contract.value_mappings would produce.
    assert header[:4] == ("Material (Original)", "PRDID (Paired)", "Plant (Original)", "LOCID (Paired)")

    rows = {r[1]: r for r in ws.iter_rows(values_only=True, min_row=2)}  # keyed by PRDID (Paired)

    match = rows["PA"]
    assert match[0] == "A"  # Material (Original)
    assert match[2] == "P1" and match[3] == "LOC1"
    qty_col = header.index("ReqQty")
    tgt_col = header.index("SalesOrderRequest")
    delta_col = header.index("Delta")
    assert match[qty_col] == 10 and match[tgt_col] == 10 and match[delta_col] == 0

    mismatch = rows["B"]
    assert mismatch[0] == "B"
    assert mismatch[qty_col] == 20 and mismatch[tgt_col] == 25 and mismatch[delta_col] == -5

    # Traceability columns carry the batch's real ids (Auto mode, unlike
    # Manual mode's always-blank Batch ID/Record ID).
    batch_col = header.index("Batch ID")
    record_col = header.index("Record ID")
    for row in rows.values():
        assert row[batch_col] == "batch_1"
        assert row[record_col] is not None


def test_streaming_run_mapping_details_sheet_is_populated():
    run_id = _streaming_run()
    content = service.build_comparison_workbook(run_id)
    wb = openpyxl.load_workbook(BytesIO(content))

    ws = wb["Mapping Details"]
    rows = list(ws.iter_rows(values_only=True))
    data = rows[1:]
    assert data, "Mapping Details must not be empty for an Auto-mode run"

    by_source = {(r[0], r[1]): r for r in data}
    a = by_source[("Material → PRDID", "A")]
    assert a[2] == "PA" and a[3] == "Paired" and a[4] == "Verified"
    p1 = by_source[("Plant → LOCID", "P1")]
    assert p1[2] == "LOC1" and p1[3] == "Paired"


def test_streaming_run_insights_and_pdf_work_for_chatbot_view_insights():
    """The chatbot's "View Insights" pill (AssistantBot.jsx's
    handleViewInsights) hits POST /insights/pdf for exactly this run shape —
    a chat-completed Auto-mode run, never a Manual-mode run with a real
    shadow/snapshot pair. Both the insights builder (the same one the
    run-linked Insights page uses) and the PDF renderer must handle it
    without needing snapshot/shadow data."""
    run_id = _streaming_run()
    payload = insights_builder.build_for_run(run_id)

    assert payload["runId"] == run_id
    assert payload["breakRate"]["total"] == 2
    by_key = {r["key"]: r for r in payload["breakRate"]["results"]}
    assert by_key["match"]["count"] == 1
    assert by_key["quantity_mismatch"]["count"] == 1

    pdf_bytes = insights_builder.pdf_for(run_id=run_id)
    assert pdf_bytes[:4] == b"%PDF"
